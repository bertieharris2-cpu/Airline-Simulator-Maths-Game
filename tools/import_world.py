#!/usr/bin/env python3
"""Import the world workbook into the prototype's data.

Reads data/Airline-World-Workbook.xlsx (v4.2) and writes:
  * a JSON block <script type="application/json" id="data-workbook"> in src/opening/parts/p2_data.html
    (live rows only, as the workbook's README asks; "optional" rows are kept and flagged; note rows with no status are skipped);
  * data/Import-Report.md: what was loaded, what was skipped, every check that failed, and the rows dated inside
    chapter 1 that are not live (the game reports them rather than inventing values).

The opening prototype uses data-workbook: src/opening/parts/p4a2_workbook.js applies its values when the page loads, and
src/opening/engine/world-engine.js holds the rules. Rebuild after importing: python3 src/opening/build.py

Usage:
  python3 tools/import_world.py            # import and write the report
  python3 tools/import_world.py --check    # write the report only; the game data is not touched
  python3 tools/import_world.py --all      # also load placeholder and later rows (each keeps its status)

Ids never change. If a later workbook drops an id that an earlier import had, the report lists it as an error.
"""
import datetime as dt, json, re, sys
from collections import Counter
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit('Needs openpyxl: pip install openpyxl')

ROOT = Path(__file__).resolve().parent.parent
BOOK = ROOT / 'data' / 'Airline-World-Workbook.xlsx'
DATA = ROOT / 'src' / 'opening' / 'parts' / 'p2_data.html'
REPORT = ROOT / 'data' / 'Import-Report.md'
PASS1 = ROOT / 'data' / 'fable-routes-pass1.json'     # the engine rules (Fable Pass 1, section 0): bands, duty padding
BLOCK_ID = 'data-workbook'

# The game already uses these ids for the same aircraft.
AIRCRAFT_IDS = {'saab340': 'sf34', 'atr72': 'at72'}
BANDS = ['early', 'midmorning', 'midday', 'afternoon', 'evening']
HEADER_NAMES = {'#': 'n', 'leg hours TO / Saab-ATR / jet': 'legHours'}

args = set(sys.argv[1:])
LOAD = {'live', 'optional'} | ({'placeholder', 'later'} if '--all' in args else set())
issues = []          # (level, sheet, message): level is 'error' (fix before wiring in), 'check' (probably wrong) or 'note'
def issue(level, sheet, msg): issues.append((level, sheet, msg))


# ---------- reading ----------
def key(h):
    h = str(h).strip()
    if h in HEADER_NAMES: return HEADER_NAMES[h]
    words = re.findall(r'[A-Za-z0-9]+', re.sub(r'\(.*?\)', '', h).replace('£/L', ''))
    return words[0][0].lower() + words[0][1:] + ''.join(w[0].upper() + w[1:] for w in words[1:])

def plain(v):
    if isinstance(v, (dt.datetime, dt.date)): return v.strftime('%Y-%m-%d')
    if isinstance(v, str):
        v = v.strip()
        if v.upper() in ('TRUE', 'FALSE'): return v.upper() == 'TRUE'
        return v or None
    if isinstance(v, float) and v.is_integer(): return int(v)
    return v

def sheet(wb, name):
    ws = wb[name]; rows = list(ws.iter_rows(values_only=True))
    head = [key(h) for h in rows[0] if h is not None]
    out = []
    for r in rows[1:]:
        if all(c is None for c in r): continue
        out.append({k: plain(v) for k, v in zip(head, r)})
    return out

counts = {}
def live(name, rows):
    st = Counter((r.get('status') or 'note row') for r in rows)
    keep = [r for r in rows if r.get('status') in LOAD]          # a row with no status is a note under the table
    counts[name] = (len(keep), dict(st))
    return keep


# ---------- small parsers ----------
def num(s):
    s = str(s).strip().replace('£', '').replace(',', '')
    f = float(s); return int(f) if f.is_integer() else f

def nums(s):            # '70|80|90' -> [70, 80, 90]
    return [num(x) for x in str(s).split('|') if x.strip()] if s not in (None, '') else []

def pairs(s):           # '70:80|80:72' -> {'70': 80, '80': 72}
    if not s or ':' not in str(s): return None
    return {str(num(a)): num(b) for a, b in (p.split(':') for p in str(s).split('|') if p.strip())}

def pct(s):             # '+15%' -> 0.15
    return None if s is None else round(num(str(s).replace('%', '').replace('+', '')) / 100, 4)

def changes(s):         # 'business:-30%,gva:+10' -> {'business': {'pct': -30}, 'gva': {'add': 10}}
    if not s: return None
    out = {}
    for part in str(s).split(','):
        if ':' not in part:                      # free text such as 'weekend multipliers (Settings)': the game reads Settings for that
            out['note'] = part.strip(); continue
        k, v = part.split(':'); v = v.strip()
        out[k.strip()] = {'pct': num(v.rstrip('%').replace('+', ''))} if v.endswith('%') else {'add': num(v.replace('+', ''))}
    return out

def fares(s):           # 'dub:50' -> {'dub': 50}
    return {k.strip(): num(v) for k, v in (p.split(':') for p in str(s).split(','))} if s else None

def words(s):           # 'fuel,pricing' -> ['fuel', 'pricing']
    return [w.strip() for w in str(s).split(',') if w.strip()] if s else []

def unlocks(s):         # 'routes:dub,par; planes:dhc6(launch deal); mechanics:fare' -> {'routes': [...], ...}
    if not s: return None
    out = {}
    for part in str(s).split(';'):
        if ':' not in part: continue
        k, v = part.split(':', 1)
        items = [re.sub(r'\(.*?\)', '', x).strip() for x in v.split(',')]
        if k.strip() == 'planes': items = [AIRCRAFT_IDS.get(x, x) for x in items]
        out[k.strip()] = [x for x in items if x]
    return out

def date(s): return dt.date.fromisoformat(str(s)[:10])


# ---------- the workbook ----------
wb = openpyxl.load_workbook(BOOK, data_only=True)

settings_rows = sheet(wb, 'Settings')
settings = {r['key']: r['value'] for r in live('Settings', settings_rows)}
settings_status = {r['key']: (r.get('status') or 'note row') for r in settings_rows if r.get('key')}
start = date(settings['startDate'])
if 'blockedFlightCodes' in settings: settings['blockedFlightCodes'] = [w.upper() for w in words(settings['blockedFlightCodes'])]
if 'crewFromDate' in settings: settings['crewFromDate'] = date(settings['crewFromDate']).isoformat()   # crew duty waits for the airline-type choice (week-1 review)
CH1_END = dt.date(2030, 12, 28)   # chapter 1's review date; rows dated on or before it must be live (see "Chapter 1 touches")

arche = live('Archetypes', sheet(wb, 'Archetypes'))
archetypes = []
for a in arche:
    archetypes.append({'id': a['id'], 'name': a['name'], 'bands': {b: a[b] for b in BANDS},
                       'timeSensitiveShare': a['timeSensitiveShare'], 'fareSensitivity': a['fareSensitivity'],
                       'businessShare': a['businessShare'], 'growthPerWeek': a['growthPerWeek'], 'growthCeiling': a['growthCeiling'],
                       'season': {s: a['season' + s.title()] for s in ('spring', 'summer', 'autumn', 'winter')},
                       'story': a['story'], 'status': a['status']})

all_routes = sheet(wb, 'Routes')
routes = []
for r in live('Routes', all_routes):
    routes.append({'id': r['id'], 'city': r['city'], 'country': r['country'], 'airport': r['airportCode'],
                   'lat': r['lat'], 'lon': r['lon'], 'utc': r['utc'], 'archetype': r['archetype'], 'km': r['distanceKm'],
                   'turnaroundAwayMin': r['turnaroundAwayMin'], 'landingFeeAway': r['landingFeeAway'],
                   'baseFare': r['baseFare'], 'fareOptions': nums(r['fareOptions']), 'demandAtFare': pairs(r['demandAtFare']),
                   'competitionSensitivity': r['competitionSensitivity'], 'unlockDay': r['unlockDay'],
                   'story': r['story'], 'status': r['status']})

catalogue = [{'id': r['id'], 'city': r['city'], 'country': r['country'], 'airport': r['code'], 'km': r['distanceKm'],
              'archetype': r['archetype'], 'haul': r['haul'], 'baseFare': r['baseFare'], 'legHours': r['legHours'],
              'plannedOpenWeek': r['plannedOpenWeek'], 'notes': r['notes'], 'status': r['status']}
             for r in live('RouteCatalogue', sheet(wb, 'RouteCatalogue'))]

all_aircraft = sheet(wb, 'Aircraft')
aircraft = [{'id': AIRCRAFT_IDS.get(a['id'], a['id']), 'wbId': a['id'], 'name': a['name'], 'tier': a['tier'], 'seats': a['seats'],
             'speedKmh': a['speedKmh'], 'rangeKm': a['rangeKm'], 'fuelPer100Km': a['fuelPer100Km'], 'hourlyCost': a['hourlyCost'],
             'dayCost': a['dayCost'], 'listPrice': a['listPrice'], 'launchDealPrice': a['launchDealPrice'],
             'shopFromDay': a['shopFromDay'], 'identity': a['identity'], 'status': a['status'],
             'artworkId': a.get('artworkId'), 'photoFile': a.get('photoFile'), 'photoCredit': a.get('photoCredit'), 'photoLicence': a.get('photoLicence'), 'photoSource': a.get('photoSource')}
            for a in live('Aircraft', all_aircraft)]
# The showroom's "coming soon" row (CR5): the aircraft not yet live, by name, date and picture only. No seats, prices or costs are read.
aircraft_preview = [{'id': AIRCRAFT_IDS.get(a['id'], a['id']), 'name': a['name'], 'tier': a['tier'], 'shopFromDay': a['shopFromDay'], 'status': a['status'],
                     'photoFile': a.get('photoFile'), 'photoCredit': a.get('photoCredit'), 'photoLicence': a.get('photoLicence'), 'photoSource': a.get('photoSource')}
                    for a in all_aircraft if a.get('id') and a.get('status') in ('placeholder', 'later')]

finance_rows = live('Finance', sheet(wb, 'Finance'))
TUTORIAL = {'dhc6', 'saab340', 'atr72', 'e175', 'e190'}      # launch-deal finance is paid by the day; from the A220 by the week
def fin_num(v): return v if isinstance(v, (int, float)) else None
finance = [{'aircraft': AIRCRAFT_IDS.get(f['aircraftId'], f['aircraftId']), 'cashPrice': f['cashPrice'], 'leaseDaily': f['leaseDaily'],
            'financeDeposit': f['financeDeposit'], 'financePayment': f['financeWeeklyPayment'], 'financeCount': f['financeWeeks'],
            'paymentUnit': 'day' if f['aircraftId'] in TUTORIAL else 'week', 'financeTotalPaid': fin_num(f['financeTotalPaid']),
            'leaseFrom': f['leaseFrom'], 'buyFrom': f['buyFrom'], 'financeFrom': f['financeFrom'], 'notes': f['notes'], 'status': f['status']}
           for f in finance_rows]

airports = [{'code': a['code'], 'name': a['name'], 'terminal': a['terminal'], 'landingFee': a['landingFee'],
             'demandModifier': pct(a['demandModifier']), 'tendency': a['tendency'], 'turnaroundMin': a['turnaroundMin'],
             'passengerCharge': a['passengerCharge'], 'notes': a['notes'], 'status': a['status']}
            for a in live('Airports', sheet(wb, 'Airports'))]

all_cal = sheet(wb, 'Calendar')
calendar = [{'day': c['day'], 'date': c['date'], 'weekday': c['weekday'], 'week': c['week'], 'phase': c['phase'],
             'fuel': c['fuel'], 'title': c['title'], 'headlines': [h for h in (c['headline1'], c['headline2']) if h],
             'event': c['eventId'], 'unlocks': unlocks(c['unlocks']), 'demandChanges': changes(c['demandChanges']),
             'rivalFares': fares(c['rivalFares']), 'jobs': words(c['jobs']), 'gatedCalc': c['gatedCalc'],
             'challenge': c['challenge'], 'status': c['status']}
            for c in live('Calendar', all_cal)]

all_market = sheet(wb, 'Market')
market = [{'week': m['week'], 'weekStart': m['weekStart'], 'season': m['season'], 'fuel': m['fuel'],
           'businessMult': m['businessMult'], 'leisureMult': m['leisureMult'],
           'seasonMult': {k[:-len('SeasonMult')]: v for k, v in m.items() if k.endswith('SeasonMult')},
           'weather': m['weather'], 'headline': m['headline'], 'status': m['status']}
          for m in live('Market', all_market)]

all_events = sheet(wb, 'Events')
ev_rows = live('Events', all_events)
events, by_id = [], {}
for e in ev_rows:
    x = by_id.get(e['eventId'])
    if not x:
        x = by_id[e['eventId']] = {'id': e['eventId'], 'date': e['date'], 'title': e['title'], 'text': e['text'],
                                   'options': [], 'effects': None, 'costShare': e.get('costShare'), 'status': e['status']}
        events.append(x)
    if e['option'] is not None:
        x['options'].append({'n': e['option'], 'label': e['label'], 'sub': e['subLabel'], 'cash': e['cash'],
                             'stars': e['stars'], 'board': e['boardStatus'], 'effects': e['otherEffects'], 'why': e['whyLine']})
    elif e['otherEffects']:
        x['effects'] = e['otherEffects']

all_ch = sheet(wb, 'Challenges')
challenges = [{'n': c['n'], 'fromDate': c.get('fromDate'), 'title': c['title'], 'story': c['story'], 'question': c['question'],
               'answer': c['answer'], 'unit': c['unit'], 'reward': c['reward'], 'status': c['status']}
              for c in live('Challenges', all_ch)]

mechanics = [{'fromDay': m['fromDay'], 'mechanic': m['mechanic'], 'notes': m['notes'], 'status': m['status']}
             for m in live('Mechanics', sheet(wb, 'Mechanics'))]

# v4 sheets: chapters pace the game, hand-sum rules decide what is done by hand, catering replaces the game's snack options, hunts are read only
def opt_sheet(name): return sheet(wb, name) if name in wb.sheetnames else []
if any(n not in wb.sheetnames for n in ('Chapters', 'HandSumRules', 'Catering', 'Hunts')):
    issue('error', 'Workbook', 'a v4 sheet is missing: ' + ', '.join(n for n in ('Chapters', 'HandSumRules', 'Catering', 'Hunts') if n not in wb.sheetnames))
all_chapters = opt_sheet('Chapters')
chapters = [{'chapter': c['chapter'], 'name': c['name'], 'start': c['start'], 'end': c['end'], 'setUpDays': c['setUpDays'], 'runCadence': c['runCadence'],
             'reviewDate': c['reviewDate'], 'newIdea': c['newIdea'], 'mechanicsArriving': c['mechanicsArriving'], 'handSumsPlanned': c['handSumsPlanned'],
             'mathsFront': c['mathsFront'], 'status': c['status']} for c in live('Chapters', all_chapters)]
handSumRules = [{'rule': r['rule'], 'name': r['name'], 'text': r['ruleText'], 'trigger': r['trigger'], 'examples': r['examples'], 'status': r['status']}
                for r in live('HandSumRules', opt_sheet('HandSumRules'))]
catering = [{'id': c['optionId'], 'name': c['name'], 'sellingPrice': c['sellingPrice'], 'takeUp': c['takeUp'], 'unitStockCost': c['unitStockCost'],
             'freeCostPerPassenger': c['freeCostPerPassenger'], 'reputationEffect': c['reputationEffect'], 'notes': c['notes'], 'status': c['status']}
            for c in live('Catering', opt_sheet('Catering'))]
all_hunts = opt_sheet('Hunts')
# Livery (v4.3 draft, CR5): the paint shop's fixed choices, grouped by kind
all_livery = opt_sheet('Livery')
if 'Livery' not in wb.sheetnames: issue('error', 'Livery', 'no Livery sheet: the paint shop has no palette, stripes, tail symbols, logo shapes or name suggestions.')
livery = {'colours': [], 'stripes': [], 'tails': [], 'logoShapes': [], 'planeNames': []}
LIVERY_KINDS = {'colour': 'colours', 'stripe': 'stripes', 'tail': 'tails', 'logoShape': 'logoShapes', 'planeName': 'planeNames'}
for r in live('Livery', all_livery):
    k = LIVERY_KINDS.get(r.get('kind'))
    if not k: issue('check', 'Livery', f'row {r.get("id")}: unknown kind `{r.get("kind")}` (colour, stripe, tail, logoShape or planeName).'); continue
    row = {'id': r['id'], 'name': r['name']}
    if k == 'colours':
        if not re.fullmatch(r'#[0-9A-Fa-f]{6}', str(r.get('value') or '')): issue('check', 'Livery', f'colour {r["id"]} has no hex value.'); continue
        row['hex'] = str(r['value']).upper()
    livery[k].append(row)
for k, least in (('colours', 6), ('stripes', 2), ('tails', 2), ('logoShapes', 2), ('planeNames', 4)):
    if len(livery[k]) < least: issue('check', 'Livery', f'only {len(livery[k])} live {k}: the paint shop wants at least {least}.')
hunts = [{'id': h['huntId'], 'date': h['date'], 'chapter': h['chapter'], 'host': h['host'], 'whatIsWrong': h['whatIsWrong'], 'pupilChecks': h['pupilChecks'],
          'consequenceIfMissed': h['consequenceIfMissed'], 'status': h['status']} for h in live('Hunts', all_hunts)]

# The engine rules the workbook assumes. Band times and duty padding are not on a sheet yet, so they come from Pass 1.
engine = {}
if PASS1.exists():
    p1 = json.loads(PASS1.read_text())
    a = p1.get('assumptions', {})
    engine = {k: a[k] for k in ('serviceIsRoundTrip', 'revenuePerService', 'bands', 'allocation', 'dutyPaddingMinutes', 'airportOpen') if k in a}
    engine['source'] = PASS1.name
    for k, sk in (('secondCrewCost', 'secondCrewCost'), ('maxDutyHours', 'maxDutyHours')):
        if k in a and sk in settings and a[k] != settings[sk]:
            issue('check', 'Settings', f'{sk} is {settings[sk]}, but Pass 1 has {a[k]}. The workbook wins.')
    p1a = {x['id']: x for x in p1.get('aircraft', [])}
    for x in all_aircraft:
        y = p1a.get(x['id'])
        if not y or x['status'] != 'live': continue
        diff = [f'{k} {y[k]} → {x[k]}' for k in ('seats', 'speedKmh', 'fuelPer100Km', 'hourlyCost', 'dayCost', 'listPrice') if k in y and y[k] != x[k]]
        if diff: issue('note', 'Aircraft', f'{x["id"]} differs from Pass 1 (Pass 1 → workbook): {", ".join(diff)}. The workbook wins.')
else:
    issue('error', 'Settings', f'{PASS1.name} is missing, so the day bands and duty padding are unknown.')

world = {'source': BOOK.name, 'engine': engine, 'imported': dt.date.today().isoformat(), 'loaded': sorted(LOAD),
         'settings': settings, 'archetypes': archetypes, 'routes': routes, 'routeCatalogue': catalogue,
         'aircraft': aircraft, 'finance': finance, 'airports': airports, 'calendar': calendar, 'market': market,
         'events': events, 'challenges': challenges, 'mechanics': mechanics,
         'chapters': chapters, 'handSumRules': handSumRules, 'catering': catering, 'hunts': hunts,
         'livery': livery, 'aircraftPreview': aircraft_preview}


# ---------- checks ----------
WD = lambda d: d.strftime('%a')

# dates and weekdays
note = next((r.get('notes') or '' for r in settings_rows if r['key'] == 'startDate'), '')
import re
first_wd = re.search(r'Monday|Tuesday|Wednesday|Thursday|Friday|Saturday|Sunday', note)
for wd in ([first_wd.group(0)] if first_wd else []):
    if start.strftime('%A') != wd:
        nxt = start + dt.timedelta(days=(['Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday', 'Saturday', 'Sunday'].index(wd) - start.weekday()) % 7)
        issue('error', 'Settings', f'startDate {start:%d %b %Y} is a {start:%A}, but its note says "{wd}". '
              f'The nearest {wd} after it is {nxt:%d %b %Y}.')
for c in all_cal:
    d = date(c['date'])
    if d != start + dt.timedelta(days=c['day'] - 1):
        issue('error', 'Calendar', f'day {c["day"]} is dated {d}, but startDate + {c["day"] - 1} days is {start + dt.timedelta(days=c["day"] - 1)}.')
    if c['weekday'] and c['weekday'] != WD(d):
        issue('error', 'Calendar', f'day {c["day"]} ({d}) is labelled {c["weekday"]}, but it is a {WD(d)}.')
    for wd in ('Saturday', 'Sunday', 'Monday'):
        if (c['title'] or '').startswith(wd) and d.strftime('%A') != wd:
            issue('error', 'Calendar', f'day {c["day"]} is titled "{c["title"]}", but {d:%d %b} is a {d:%A}.')
    if c['phase'] == 'weekend' and d.weekday() < 5:
        issue('error', 'Calendar', f'day {c["day"]} is a weekend row, but {d:%d %b} is a {d:%A}.')
for m in all_market:
    want = start + dt.timedelta(days=7 * (m['week'] - 1))
    if date(m['weekStart']) != want:
        issue('error', 'Market', f'week {m["week"]} starts {m["weekStart"]}, expected {want}.'); break
weeks = [m['week'] for m in all_market]
if weeks != list(range(1, len(weeks) + 1)): issue('error', 'Market', 'week numbers are not 1, 2, 3… in order.')
if all_market and date(all_market[0]['weekStart']).weekday() != 0:
    issue('check', 'Market', f'weeks start on a {date(all_market[0]["weekStart"]):%A}, not a Monday, so a Mon–Fri game week straddles two Market rows.')

# fuel: the Calendar's daily price against the Market's weekly one
mk = {m['week']: m for m in all_market}
for wk in sorted({c['week'] for c in all_cal}):
    days = [c for c in all_cal if c['week'] == wk and c['fuel'] is not None]
    if wk in mk and days:
        lo, hi, w = min(c['fuel'] for c in days), max(c['fuel'] for c in days), mk[wk]['fuel']
        if not (lo <= w <= hi):
            span = f'£{lo:.2f}' if lo == hi else f'£{lo:.2f}–£{hi:.2f}'
            issue('check', 'Market', f'week {wk} fuel is £{w:.2f}, but the Calendar has {span} on its days. '
                  'Time skips use the Market price, so the two should agree.')

# archetypes
aids = {a['id'] for a in arche}
for a in arche:
    tot = round(sum(a[b] for b in BANDS), 6)
    if abs(tot - 1) > 1e-6: issue('error', 'Archetypes', f'{a["id"]}: day-band shares add up to {tot}, not 1.')
for r in all_routes + sheet(wb, 'RouteCatalogue'):
    if r.get('archetype') and r['archetype'] not in {a['id'] for a in sheet(wb, 'Archetypes')}:
        issue('error', 'Routes', f'{r["id"]}: archetype "{r["archetype"]}" is not on the Archetypes sheet.')
    elif r.get('archetype') and r['archetype'] not in aids and (r.get('status') in LOAD):
        issue('check', 'Routes', f'{r["id"]} is loaded, but its archetype "{r["archetype"]}" is not live. The game uses the route\'s own demand table and applies no business/leisure or season multipliers to it.')

# routes: demand tables
for r in routes:
    t = r['demandAtFare']
    if not t: issue('error', 'Routes', f'{r["id"]} is live but has no demand table.'); continue
    fs = sorted(t, key=float)
    if any(t[a] < t[b] for a, b in zip(fs, fs[1:])):
        issue('check', 'Routes', f'{r["id"]}: demand rises with the fare somewhere in {r["demandAtFare"]}.')
    miss = [f for f in r['fareOptions'] if str(f) not in t]
    if miss: issue('error', 'Routes', f'{r["id"]}: fare options {miss} have no demand figure.')
    if r['baseFare'] not in r['fareOptions']: issue('check', 'Routes', f'{r["id"]}: base fare £{r["baseFare"]} is not one of the fare options.')

# route stories against the distances (the pupil sees both side by side when choosing a first market)
for grp in sorted({r['unlockDay'] for r in routes}):
    rs = [r for r in routes if r['unlockDay'] == grp]
    if len(rs) < 2: continue
    lo, hi = min(r['km'] for r in rs), max(r['km'] for r in rs)
    both = ', '.join(f"{x['city']} {x['km']} km" for x in rs)
    for r in rs:
        t = (r['story'] or '').lower()
        if re.search(r'\bshort\b', t) and r['km'] > lo:
            issue('error', 'Routes', f'{r["id"]}: the story says "short", but it is the longer of the routes opening on day {grp} ({both}). The game shows the story next to the flight time.')
        if re.search(r'\blong(er)?\b', t) and r['km'] < hi:
            issue('error', 'Routes', f'{r["id"]}: the story says "longer", but it is the shorter of the routes opening on day {grp} ({both}). The game shows the story next to the flight time.')

# leg times on quarter hours (the README sheet's promise)
live_planes = [a for a in all_aircraft if a['status'] == 'live']
for r in [x for x in all_routes if x['status'] == 'live']:
    bad = []
    for a in live_planes:
        mins = r['distanceKm'] / a['speedKmh'] * 60
        if abs(mins / 15 - round(mins / 15)) > 1e-9: bad.append(f'{a["name"]} {mins:.0f} min')
    if bad: issue('check', 'Routes', f'{r["id"]} ({r["distanceKm"]} km): legs not on a quarter hour: {", ".join(bad)}.')

# references between sheets
ev_ids = {e['eventId'] for e in all_events}
ch_by_n = {c['n']: c for c in all_ch}
route_ids = {r['id'] for r in all_routes}
plane_ids = {AIRCRAFT_IDS.get(a['id'], a['id']) for a in all_aircraft}
for c in all_cal:
    if c['eventId'] and c['eventId'] not in ev_ids: issue('error', 'Calendar', f'day {c["day"]}: event "{c["eventId"]}" is not on the Events sheet.')
    if c['eventId'] in ev_ids:
        ed = next(e['date'] for e in all_events if e['eventId'] == c['eventId'])
        if ed != c['date']: issue('error', 'Events', f'"{c["eventId"]}" is dated {ed}, but the Calendar puts it on day {c["day"]} ({c["date"]}).')
    if c['challenge'] is not None:
        ch = ch_by_n.get(c['challenge'])
        if not ch: issue('error', 'Calendar', f'day {c["day"]}: challenge {c["challenge"]} is not on the Challenges sheet.')
        else: issue('note', 'Calendar', f'day {c["day"]} names challenge {c["challenge"]}, but Settings captainsChallengeCadence starts challenges on {ch.get("fromDate")} (one a term). The game follows the Challenges sheet dates; the Calendar column is ignored.')
    u = unlocks(c['unlocks']) or {}
    for rid in u.get('routes', []):
        if rid not in route_ids and not (' ' in rid):      # 'first market (pupil chooses dub or par)' is a description, not an id
            issue('error', 'Calendar', f'day {c["day"]} unlocks route "{rid}", which is not on the Routes sheet.')
    for pid in u.get('planes', []):
        if pid not in plane_ids: issue('error', 'Calendar', f'day {c["day"]} unlocks plane "{pid}", which is not on the Aircraft sheet.')
for f in finance_rows:
    if AIRCRAFT_IDS.get(f['aircraftId'], f['aircraftId']) not in plane_ids: issue('error', 'Finance', f'"{f["aircraftId"]}" is not on the Aircraft sheet.')
    tot = fin_num(f['financeTotalPaid'])
    if tot is not None and f['financeDeposit'] + f['financeWeeklyPayment'] * f['financeWeeks'] != tot:
        issue('check', 'Finance', f'{f["aircraftId"]}: deposit + payment × {f["financeWeeks"]} = {f["financeDeposit"] + f["financeWeeklyPayment"] * f["financeWeeks"]:,}, but financeTotalPaid says {tot:,} (rounded to £1,000 as the note says).')
    if f['leaseDaily'] and f['cashPrice'] and f['aircraftId'] not in TUTORIAL and f['cashPrice'] / f['leaseDaily'] != 1000:
        issue('check', 'Finance', f'{f["aircraftId"]}: price ÷ lease a day = {f["cashPrice"] / f["leaseDaily"]:g}, not the 1,000 days of F4.')
dates = [e['date'] for e in events]
if dates != sorted(dates): issue('note', 'Events', 'events are not in date order (harmless; the game sorts them).')

# reputation is switched off in the prototype
star_ev = [e['id'] for e in events if any(o['stars'] for o in e['options'])]
star_ch = [str(c['n']) for c in challenges if 'star' in str(c['reward'])]
if star_ev or star_ch:
    issue('note', 'Events', f'reputation is switched off in the prototype, so star effects do nothing yet: events {", ".join(star_ev) or "none"}; '
          f'challenges {", ".join(star_ch) or "none"}. Those choices then differ only in cash.')

# ids never change: compare with the last import
src = DATA.read_text()
m = re.search(r'<script type="application/json" id="%s">(.*?)</script>' % BLOCK_ID, src, re.S)
if m:
    old = json.loads(m.group(1))
    removed = {AIRCRAFT_IDS.get(x.get('id'), x.get('id')) for x in all_aircraft + all_routes + all_events if x.get('status') == 'removed'}
    for part, idk in (('routes', 'id'), ('aircraft', 'id'), ('archetypes', 'id'), ('events', 'id')):
        gone = {x[idk] for x in old.get(part, [])} - {x[idk] for x in world[part]}
        if gone & removed: issue('note', part.title(), f'now marked removed and no longer loaded: {", ".join(sorted(gone & removed))}. The id is kept on the sheet.')
        if gone - removed: issue('error', part.title(), f'ids in the last import are missing now: {", ".join(sorted(gone - removed))}. Ids never change.')


# ---------- what changes when this is wired in ----------
def J(i):
    mm = re.search(r'<script type="application/json" id="%s">(.*?)</script>' % i, src, re.S)
    return json.loads(mm.group(1)) if mm else None
GW, GP = J('data-world') or {}, {p['id']: p for p in (J('data-planes') or [])}
rows = []
for r in routes:
    g = next((x for x in GW.get('routes', []) if x['id'] == r['id']), None)
    if g: rows.append(f'| {r["city"]} distance | {g["km"]} km | {r["km"]} km |')
    if g: rows.append(f'| {r["city"]} fares | {", ".join("£%d" % p for p in g["prices"])} | {", ".join("£%d" % p for p in r["fareOptions"])} |')
for a in aircraft:
    g = GP.get(a['id'])
    if not g: continue
    rows.append(f'| {a["name"]} price | £{g["price"]:,} | £{a["listPrice"]:,}' + (f' (launch deal £{a["launchDealPrice"]:,})' if a['launchDealPrice'] else '') + ' |')
    if g.get('hourCost') or g.get('dayCost'):
        rows.append(f'| {a["name"]} running costs | £{g.get("hourCost", "—")} an hour, £{g.get("dayCost", "—")} a day | £{a["hourlyCost"]} an hour, £{a["dayCost"]} a day |')
    rows.append(f'| {a["name"]} fuel | {g["fuelUse"]} L per 100 km | {a["fuelPer100Km"]} L per 100 km |')
rows.append(f'| Second crew | £{GW.get("crewCost", "—")} | £{settings.get("secondCrewCost")} when duty is over {settings.get("maxDutyHours")} h, from {settings.get("crewFromDate", "the start")} |')
rows.append(f'| Starting cash | £{GW.get("startingCash", "—"):,} | £{settings.get("startingCash"):,} |')
rows.append(f'| Start date | {GW.get("startDate", "Sun 12 May 2030")} | {start:%a %d %b %Y} |')


# ---------- write ----------
order = {'error': 0, 'check': 1, 'note': 2}
issues.sort(key=lambda x: (order[x[0]], x[1]))
n = Counter(l for l, _, _ in issues)
out = [f'# Import report: {BOOK.name}', '',
       f'*Written by `tools/import_world.py` on {dt.date.today():%d %b %Y}. Loaded: {", ".join(sorted(LOAD))} rows. '
       f'{"Check only: the game data was not changed." if "--check" in args else f"Written to `{DATA.relative_to(ROOT)}` as `{BLOCK_ID}`; rebuild the prototype to use it."}*', '',
       f'**{n["error"]} to fix, {n["check"]} to check, {n["note"]} notes.**', '',
       '## Loaded', '', '| Sheet | Loaded | Rows by status |', '| --- | --- | --- |']
for s, (k, st) in counts.items():
    k = f'{k} rows → {len(events)} events' if s == 'Events' else k   # one row per option: the game holds one event per id
    out.append(f'| {s} | {k} | {", ".join(f"{v} {s2}" for s2, v in sorted(st.items()))} |')
for title, lvl in (('To fix before wiring in', 'error'), ('To check', 'check'), ('Notes', 'note')):
    xs = [(s, msg) for l, s, msg in issues if l == lvl]
    if xs: out += ['', f'## {title}', ''] + [f'- **{s}:** {msg}' for s, msg in xs]
# chapter 1 touches: every row dated inside chapter 1 that the game would read but is not live, and anything it needs that is missing
touch = []
def dated(rows, when, label):
    for r in rows:
        d = r.get(when)
        try: dd = date(d) if d else None
        except Exception: dd = None
        if dd and dd <= CH1_END and r.get('status') not in LOAD: touch.append(f'- **{label}:** {r.get("eventId") or r.get("id") or r.get("huntId") or r.get("n") or r.get("week") or r.get("day")} on {dd:%d %b %Y} is `{r.get("status")}`: not loaded.')
dated(all_cal, 'date', 'Calendar'); dated(all_market, 'weekStart', 'Market'); dated([e for e in all_events if e.get('title')], 'date', 'Events')
dated(all_ch, 'fromDate', 'Challenges'); dated(all_hunts, 'date', 'Hunts')
ch1_days = 118   # 2 Sep to 28 Dec 2030
for r in all_routes:
    if isinstance(r.get('unlockDay'), int) and r['unlockDay'] <= ch1_days:
        if r.get('status') not in LOAD: touch.append(f'- **Routes:** {r["id"]} opens on day {r["unlockDay"]} but is `{r.get("status")}`: not loaded.')
        elif not pairs(r.get('demandAtFare')): touch.append(f'- **Routes:** {r["id"]} opens on day {r["unlockDay"]} with no demand table: it stays locked and the game says so.')
        elif r.get('archetype') not in aids: touch.append(f'- **Routes:** {r["id"]} opens on day {r["unlockDay"]} but its archetype `{r.get("archetype")}` is not live: no weekend or season multipliers on this route, and the game says so on its route card.')
for a in all_aircraft:
    if isinstance(a.get('shopFromDay'), int) and a['shopFromDay'] <= ch1_days and a.get('status') not in LOAD:
        touch.append(f'- **Aircraft:** {a["id"]} is in the shop from day {a["shopFromDay"]} but is `{a.get("status")}`: not loaded.')
for f in sheet(wb, 'Finance'):
    pid = f.get('aircraftId'); a = next((x for x in all_aircraft if x.get('id') == pid), None)
    if a and isinstance(a.get('shopFromDay'), int) and a['shopFromDay'] <= ch1_days and f.get('status') not in LOAD:
        touch.append(f'- **Finance:** {pid} is in the shop from day {a["shopFromDay"]} but its finance row is `{f.get("status")}`: no price, rent or finance terms.')
NEEDED = ['startDate', 'startingCash', 'airportOpen', 'airportClose', 'homeTurnaroundMin', 'secondCrewCost', 'maxDutyHours', 'fuelLotLitres', 'playableDays',
          'weekendBusinessMultSat', 'weekendBusinessMultSun', 'weekendLeisureMult', 'cashReserve', 'negativeProjectionWarning', 'eventCostScaleFromDay',
          'loadFactorPercentFromDay', 'fuelHandSumPriceStep', 'fuelHandSumLitresShare', 'snacksHandSumOnce', 'timeSkip', 'leaseFromDate', 'buyFromDate',
          'financeFromDate', 'reputationFromDate', 'quarterlyTaskFromDate', 'huntPerReview', 'captainsChallengeCadence', 'fareStepWeek2', 'reportCadence',
          'blockedFlightCodes', 'soundDefault', 'revealSeconds', 'crewFromDate']
for k in NEEDED:
    st_ = settings_status.get(k)
    if st_ is None: touch.append(f'- **Settings:** `{k}` is not on the sheet: the game needs it for chapter 1.')
    elif st_ not in LOAD: touch.append(f'- **Settings:** `{k}` is `{st_}`: not loaded' + (' (the game builds teacher skips of 0 / 1 / 4 / 8 weeks from the brief and reports it here).' if k == 'timeSkip' else ' (a date gate only in chapter 1).' if k == 'reputationFromDate' else '.'))
if not catering: touch.append('- **Catering:** no live rows: the game has no snack options.')
if not chapters or chapters[0].get('chapter') != 1: touch.append('- **Chapters:** chapter 1 is not live: nothing paces the game.')
out += ['', '## Chapter 1 touches', '',
        'Rows dated on or before 28 Dec 2030 (chapter 1\'s review) that the game would read but are not `live`, and values the chapter-1 flow needs that the workbook lacks. '
        'The game reports these; it never invents a value.', ''] + (touch or ['- Nothing: every row chapter 1 touches is live.'])

# identity and photos (CR5): the artwork and the showroom photo of every aircraft row, live or not
PHOTOS = ROOT / 'src' / 'opening' / 'assets' / 'photos'
CREDITS = ROOT / 'data' / 'Photo-Credits.md'
ph = ['| Aircraft | Status | Artwork | Photo | Credit and licence |', '| --- | --- | --- | --- | --- |']
credits = ['# Showroom photo credits', '', '*Written by `tools/import_world.py` from the Aircraft sheet\'s photo columns. Every photo in the showroom is openly licensed and shows the aircraft in plain or maker colours.*', '',
           '| Aircraft | File | Author | Licence | Source |', '| --- | --- | --- | --- | --- |']
for a in all_aircraft:
    if not a.get('id') or a.get('status') == 'removed': continue
    f = a.get('photoFile'); have = bool(f) and (PHOTOS / str(f)).exists()
    size = f'{(PHOTOS / str(f)).stat().st_size // 1024} KB' if have else ''
    art = a.get('artworkId') or ('—' if a.get('status') in LOAD else '')
    cred = ' · '.join(str(a.get(k)) for k in ('photoCredit', 'photoLicence') if a.get(k))
    photo = f'`{f}` {size}' if have else (f'`{f}` **missing on disk**' if f else 'none yet')
    ph.append(f'| {a["id"]} | {a.get("status")} | {art} | {photo} | {cred or "—"} |')
    if have and not (a.get('photoCredit') and a.get('photoLicence') and a.get('photoSource')): issue('check', 'Aircraft', f'{a["id"]}: the photo {f} has no credit, licence or source recorded: it must not be shown.')
    if f and not have: issue('check', 'Aircraft', f'{a["id"]}: photoFile {f} is not in {PHOTOS.relative_to(ROOT)}.')
    if have: credits.append(f'| {a["name"]} | {f} | {a.get("photoCredit") or "?"} | {a.get("photoLicence") or "?"} | {a.get("photoSource") or "?"} |')
CREDITS.write_text('\n'.join(credits) + ('\n' if len(credits) > 6 else '\n\nNo photos yet.\n'))
out += ['', '## Identity and photos (CR5)', '',
        f'Livery sheet: {len(livery["colours"])} colours, {len(livery["stripes"])} stripes, {len(livery["tails"])} tail symbols, {len(livery["logoShapes"])} logo shapes, {len(livery["planeNames"])} name suggestions. '
        'Photos live in `src/opening/assets/photos/` and are embedded by the build; credits are listed in `data/Photo-Credits.md`. '
        'The showroom\'s "coming soon" row reads the name, date and picture of aircraft that are not live (never their figures).', ''] + ph

out += ['', '## The prototype\'s own data and the workbook\'s', '',
        'The game uses the workbook\'s values. The prototype\'s own (in `data-world` and `data-planes`) only matter for routes and aircraft the workbook does not have yet.', '',
        '| | Prototype\'s own data | Workbook (used) |', '| --- | --- | --- |'] + rows
out += ['', '## Engine rules', '',
        'From the README sheet and Fable Pass 1, section 0, as `src/opening/engine/world-engine.js` implements them:', '',
        '- A service is a round trip that earns one plane-load at the fare (return tickets).',
        '- Five day bands: ' + (', '.join(f'{b} {v[0]:02d}:00–{v[1]:02d}:00' for b, v in engine.get('bands', {}).items()) or 'boundaries unknown') + ' (from Pass 1; the prototype has three). A service belongs to the band it leaves home in.',
        '- Time-locked passengers fly only in their band. In each band: demand × timeSensitiveShare × the band\'s share, rounded as at school (halves up). The rest are flexible and fill any seats left that day.',
        '- People at a fare come from the route\'s demand table, not a formula.',
        '- Costs per service: flying hours × hourly cost + landing fees at both ends + fuel + a charge for each passenger at home. Per day: the day cost, plus a second crew when duty is over 12 h (30 minutes before the first departure to 30 minutes after the last arrival).',
        '- Fuel: the Calendar has a price for each day; the Market sheet has one for each week, used for time skips.', '']
REPORT.write_text('\n'.join(out))

if '--check' not in args:
    block = '<script type="application/json" id="%s">%s</script>' % (BLOCK_ID, json.dumps(world, ensure_ascii=False, separators=(',', ':')))
    if m: src = src[:m.start()] + block + src[m.end():]
    else: src = src.rstrip('\n') + '\n' + block + '\n'
    DATA.write_text(src)

print(f'{REPORT.relative_to(ROOT)}: {n["error"]} to fix, {n["check"]} to check, {n["note"]} notes'
      + ('' if '--check' in args else f'; {BLOCK_ID} written ({len(routes)} routes, {len(aircraft)} aircraft, {len(calendar)} days, {len(market)} weeks, {len(events)} events, {len(chapters)} chapters, {len(catering)} catering options, {sum(len(v) for v in livery.values())} livery rows)'))
