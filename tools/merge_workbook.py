#!/usr/bin/env python3
"""Install a new world workbook from Fable, keeping what Claude Code owns.

Fable's workbooks are built from her own copy, which does not carry the identity data Claude Code added for
Change Request 5 (the paint shop and the showroom). This script copies those parts from the workbook in
data/ into the new one, writes the result to data/Airline-World-Workbook.xlsx and prints what it carried:

  * the Livery sheet (palette, stripes, tail symbols, logo shapes, plane-name suggestions);
  * the Aircraft columns artworkId, photoFile, photoCredit, photoLicence, photoSource (matched by id);
  * the Settings rows blockedFlightCodes, soundDefault, revealSeconds.

A part the new workbook already has is left as Fable wrote it. Run tools/import_world.py afterwards.

Usage: python3 tools/merge_workbook.py <new-workbook.xlsx>
"""
import sys
from copy import copy
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit('Needs openpyxl: pip install openpyxl')

ROOT = Path(__file__).resolve().parent.parent
BOOK = ROOT / 'data' / 'Airline-World-Workbook.xlsx'
AIRCRAFT_COLS = ['artworkId', 'photoFile', 'photoCredit', 'photoLicence', 'photoSource']
SETTINGS_KEYS = ['blockedFlightCodes', 'soundDefault', 'revealSeconds']

if len(sys.argv) != 2: sys.exit(__doc__)
new_path = Path(sys.argv[1])
old = openpyxl.load_workbook(BOOK)
new = openpyxl.load_workbook(new_path)
carried = []

def header(ws): return [c.value for c in ws[1]]

def copy_cell(src, dst):
    dst.value = src.value
    if src.has_style:
        dst.font, dst.fill, dst.border = copy(src.font), copy(src.fill), copy(src.border)
        dst.alignment, dst.number_format = copy(src.alignment), src.number_format

# 1. the Livery sheet
if 'Livery' not in new.sheetnames and 'Livery' in old.sheetnames:
    src = old['Livery']; dst = new.create_sheet('Livery')
    for row in src.iter_rows():
        for c in row: copy_cell(c, dst.cell(row=c.row, column=c.column))
    for k, d in src.column_dimensions.items(): dst.column_dimensions[k].width = d.width
    carried.append(f'Livery sheet ({src.max_row - 1} rows)')

# 2. the Aircraft identity columns, matched by id
oa, na = old['Aircraft'], new['Aircraft']
oh, nh = header(oa), header(na)
missing = [c for c in AIRCRAFT_COLS if c in oh and c not in nh]
if missing:
    status_col = nh.index('status') + 1 if 'status' in nh else None
    na.insert_cols(status_col, amount=len(missing)) if status_col else None
    at = status_col or len(nh) + 1
    for j, name in enumerate(missing):
        copy_cell(oa.cell(row=1, column=oh.index(name) + 1), na.cell(row=1, column=at + j))
    old_rows = {oa.cell(row=r, column=1).value: r for r in range(2, oa.max_row + 1)}
    n = 0
    for r in range(2, na.max_row + 1):
        rid = na.cell(row=r, column=1).value
        if rid in old_rows:
            for j, name in enumerate(missing):
                v = oa.cell(row=old_rows[rid], column=oh.index(name) + 1).value
                na.cell(row=r, column=at + j).value = v
                n += v is not None
    carried.append(f'Aircraft columns {", ".join(missing)} ({n} values)')

# 3. the identity settings
os_, ns = old['Settings'], new['Settings']
new_keys = {ns.cell(row=r, column=1).value for r in range(2, ns.max_row + 1)}
for r in range(2, os_.max_row + 1):
    k = os_.cell(row=r, column=1).value
    if k in SETTINGS_KEYS and k not in new_keys:
        at = ns.max_row + 1
        # keep note rows (no key) at the foot: put the setting before the first blank row after the keyed block
        for rr in range(2, ns.max_row + 1):
            if ns.cell(row=rr, column=1).value is None: at = rr; ns.insert_rows(rr); break
        for c in range(1, os_.max_column + 1): copy_cell(os_.cell(row=r, column=c), ns.cell(row=at, column=c))
        carried.append(f'Settings {k}')

new.save(BOOK)
print(f'Installed {new_path.name} as {BOOK.relative_to(ROOT)}.')
print('Carried from the previous workbook: ' + ('; '.join(carried) if carried else 'nothing (the new workbook has it all).'))
